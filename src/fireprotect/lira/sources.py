"""CSV, HTML and optional XLSX sources for LIRA force tables."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Sequence

from .errors import LiraDependencyError, LiraFormatError
from .types import RawTableRow


def _validate_headers(headers: Sequence[Any], *, source: str) -> list[str]:
    result = ["" if value is None else str(value).strip() for value in headers]
    if not result or not any(result):
        raise LiraFormatError(f"{source}: header row is empty")
    if any(not header for header in result):
        raise LiraFormatError(f"{source}: header names must not be empty")
    if len(result) != len(set(result)):
        raise LiraFormatError(f"{source}: duplicate header names are not supported")
    return result


def _rows_from_matrix(
    matrix: Sequence[Sequence[Any]], *, header_index: int, source: str
) -> list[RawTableRow]:
    if header_index < 0:
        raise LiraFormatError(f"{source}: header index cannot be negative")
    if len(matrix) <= header_index:
        raise LiraFormatError(f"{source}: header row {header_index + 1} does not exist")
    headers = _validate_headers(matrix[header_index], source=source)
    rows: list[RawTableRow] = []
    for row_number, row in enumerate(matrix[header_index + 1 :], header_index + 2):
        values = list(row)
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(values) > len(headers) and any(
            value is not None and str(value).strip() for value in values[len(headers) :]
        ):
            raise LiraFormatError(
                f"{source}: row {row_number} has more values than the header"
            )
        values.extend([None] * (len(headers) - len(values)))
        rows.append(RawTableRow(dict(zip(headers, values)), row_number))
    return rows


@dataclass(frozen=True, slots=True)
class CsvTableSource:
    path: str | Path
    encoding: str = "utf-8-sig"
    delimiter: str = ","
    header_row: int = 1

    def read_rows(self) -> list[RawTableRow]:
        if self.header_row < 1:
            raise LiraFormatError("CSV header_row must be one-based")
        with Path(self.path).open("r", encoding=self.encoding, newline="") as stream:
            matrix = list(csv.reader(stream, delimiter=self.delimiter))
        return _rows_from_matrix(
            matrix,
            header_index=self.header_row - 1,
            source=str(self.path),
        )


class _HtmlTableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.tables: list[list[list[str]]] = []
        self._table_depth = 0
        self._current_table: list[list[str]] | None = None
        self._current_row: list[str] | None = None
        self._cell_parts: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        if tag == "table":
            self._table_depth += 1
            if self._table_depth == 1:
                self._current_table = []
        elif self._table_depth == 1 and tag == "tr":
            self._current_row = []
        elif self._table_depth == 1 and tag in {"th", "td"} and self._current_row is not None:
            self._cell_parts = []

    def handle_data(self, data: str) -> None:
        if self._cell_parts is not None:
            self._cell_parts.append(data)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if self._table_depth == 1 and tag in {"th", "td"} and self._cell_parts is not None:
            assert self._current_row is not None
            self._current_row.append(" ".join("".join(self._cell_parts).split()))
            self._cell_parts = None
        elif self._table_depth == 1 and tag == "tr" and self._current_row is not None:
            assert self._current_table is not None
            if self._current_row:
                self._current_table.append(self._current_row)
            self._current_row = None
        elif tag == "table" and self._table_depth:
            if self._table_depth == 1:
                assert self._current_table is not None
                self.tables.append(self._current_table)
                self._current_table = None
            self._table_depth -= 1


@dataclass(frozen=True, slots=True)
class HtmlTableSource:
    path: str | Path
    encoding: str = "utf-8"
    table_index: int = 0
    header_row: int = 1

    def read_rows(self) -> list[RawTableRow]:
        if self.table_index < 0:
            raise LiraFormatError("HTML table_index cannot be negative")
        if self.header_row < 1:
            raise LiraFormatError("HTML header_row must be one-based")
        parser = _HtmlTableParser()
        parser.feed(Path(self.path).read_text(encoding=self.encoding))
        parser.close()
        if self.table_index >= len(parser.tables):
            raise LiraFormatError(
                f"{self.path}: table index {self.table_index} does not exist"
            )
        return _rows_from_matrix(
            parser.tables[self.table_index],
            header_index=self.header_row - 1,
            source=str(self.path),
        )


@dataclass(frozen=True, slots=True)
class XlsxTableSource:
    path: str | Path
    sheet_name: str | None = None
    header_row: int = 1
    data_only: bool = True

    def read_rows(self) -> list[RawTableRow]:
        if self.header_row < 1:
            raise LiraFormatError("XLSX header_row must be one-based")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:  # pragma: no cover - depends on test environment
            raise LiraDependencyError(
                "XLSX import requires the optional dependency 'openpyxl'; "
                "install it with: pip install openpyxl"
            ) from exc

        workbook = load_workbook(
            filename=Path(self.path), read_only=True, data_only=self.data_only
        )
        try:
            if self.sheet_name is None:
                worksheet = workbook.active
            elif self.sheet_name in workbook.sheetnames:
                worksheet = workbook[self.sheet_name]
            else:
                raise LiraFormatError(
                    f"{self.path}: worksheet {self.sheet_name!r} does not exist"
                )
            matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
        return _rows_from_matrix(
            matrix,
            header_index=self.header_row - 1,
            source=f"{self.path}:{worksheet.title}",
        )
