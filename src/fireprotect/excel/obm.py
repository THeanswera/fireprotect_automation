"""Verified-input export to the analysed 01_Общая ОБМ workbook layout."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
import json
from pathlib import Path
from typing import Any, Iterable
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from ..decision import RequiredFireResistanceDecision
from ..execution import ExecutionMode
from ..model import ProjectElement, Quantity, Unit
from ..rx3.profiles import normalize_profile_name
from ..normative import NormativeValidation
from ..technical import FireproofingTechnicalEntry, TechnicalDataStatus
from .mapping import ColumnBinding, WorkbookMapping
from .writer import ExcelCopyResult, file_sha256, write_mapped_copy


DATA_SHEET = "данные"
FIRST_ELEMENT_ROW = 6
LAST_ELEMENT_ROW = 49
EXPECTED_FORMULA_COUNT = 579


class ObmWorkbookExportError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ExcelCellChange:
    cell: str
    project_field: str
    before: str | int | float | None
    after: str | int | Decimal
    unit: str | None

    def as_dict(self) -> dict[str, Any]:
        return {
            "cell": self.cell,
            "project_field": self.project_field,
            "before": self.before,
            "after": str(self.after) if isinstance(self.after, Decimal) else self.after,
            "unit": self.unit,
        }


@dataclass(frozen=True, slots=True)
class ObmWorkbookExportReport:
    output_path: Path
    source_sha256: str
    output_sha256: str
    formula_count_before: int
    formula_count_after: int
    formulas_preserved: bool
    zip_integrity: bool
    openpyxl_opened: bool
    cell_changes: tuple[ExcelCellChange, ...]
    warnings: tuple[str, ...]
    json_report: Path
    markdown_report: Path
    export_kind: str
    technical_data_status: str
    template_verification_status: str
    recalculation_status: str

    def as_dict(self) -> dict[str, Any]:
        return {
            "output": str(self.output_path),
            "source_sha256": self.source_sha256,
            "output_sha256": self.output_sha256,
            "formula_count_before": self.formula_count_before,
            "formula_count_after": self.formula_count_after,
            "formulas_preserved": self.formulas_preserved,
            "zip_integrity": self.zip_integrity,
            "openpyxl_opened": self.openpyxl_opened,
            "cell_changes": [item.as_dict() for item in self.cell_changes],
            "warnings": list(self.warnings),
            "json_report": str(self.json_report),
            "markdown_report": str(self.markdown_report),
            "export_kind": self.export_kind,
            "technical_data_status": self.technical_data_status,
            "template_verification_status": self.template_verification_status,
            "recalculation_status": self.recalculation_status,
        }


def _quantity(value: Quantity, unit: Unit) -> Decimal:
    if not isinstance(value, Quantity):
        raise TypeError("Excel physical fields must be Quantity")
    return value.to(unit).value


def _formula_map(workbook: Any) -> dict[str, str]:
    return {
        f"{sheet.title}!{cell.coordinate}": cell.value
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }


def _decimal_cell(value: Any, *, cell: str) -> Decimal:
    try:
        result = Decimal(str(value))
    except Exception as exc:
        raise ObmWorkbookExportError(f"{cell} is not numeric: {value!r}") from exc
    if not result.is_finite():
        raise ObmWorkbookExportError(f"{cell} must be finite")
    return result


def _same_obm_profile(workbook_value: str, project_value: str) -> bool:
    workbook_name = normalize_profile_name(workbook_value)
    project_name = normalize_profile_name(project_value)
    if workbook_name == project_name:
        return True
    for prefix in ("i", "[", "l", "гн□"):
        if workbook_name.startswith(prefix) and workbook_name[len(prefix) :] == project_name:
            return True
    return False


def _validate_rows(
    elements: tuple[ProjectElement, ...],
    decisions: tuple[RequiredFireResistanceDecision, ...],
    formula_book: Any,
) -> None:
    expected = LAST_ELEMENT_ROW - FIRST_ELEMENT_ROW + 1
    if len(elements) != expected:
        raise ObmWorkbookExportError(
            f"This fixed template contains {expected} calculation rows; got {len(elements)} elements"
        )
    if len(decisions) != len(elements):
        raise ObmWorkbookExportError(
            "Each ProjectElement requires one RequiredFireResistanceDecision"
        )
    sheet = formula_book[DATA_SHEET]
    for offset, (element, decision) in enumerate(zip(elements, decisions)):
        row = FIRST_ELEMENT_ROW + offset
        element.require_fields(
            "mark",
            "profile_name",
            "area",
            "full_perimeter",
            "heated_perimeter",
            "length",
            "quantity",
            "heating_sides",
            "required_fire_resistance",
        )
        decision.require_engineer_confirmation()
        if decision.construction_type.casefold() != element.element_type.casefold():
            raise ObmWorkbookExportError(
                f"Row {row}: decision construction type does not match element"
            )
        project_r = _quantity(element.required_fire_resistance, Unit.MINUTE)  # type: ignore[arg-type]
        decision_r = _quantity(decision.required_fire_resistance, Unit.MINUTE)
        if project_r != decision_r:
            raise ObmWorkbookExportError(
                f"Row {row}: ProjectElement R={project_r} differs from decision R={decision_r} min"
            )
        if not _same_obm_profile(
            str(sheet[f"L{row}"].value or ""), element.profile_name or ""
        ):
            raise ObmWorkbookExportError(
                f"Row {row}: profile does not match the fixed workbook row"
            )
        workbook_area = _decimal_cell(sheet[f"R{row}"].value, cell=f"R{row}")
        project_area = _quantity(element.area, Unit.SQUARE_CENTIMETER)  # type: ignore[arg-type]
        if abs(workbook_area - project_area) > Decimal("0.001"):
            raise ObmWorkbookExportError(
                f"Row {row}: area conflicts with fixed workbook geometry"
            )
        full = _quantity(element.full_perimeter, Unit.MILLIMETER)  # type: ignore[arg-type]
        heated = _quantity(element.heated_perimeter, Unit.MILLIMETER)  # type: ignore[arg-type]
        if full != heated:
            raise ObmWorkbookExportError(
                f"Row {row}: workbook enforces Q=P and cannot represent the supplied heated perimeter"
            )
        # Q is a formula in the formula workbook; validate the deterministic
        # profile input formula from M/N/O instead of trusting a stale cache.
        b = _decimal_cell(sheet[f"M{row}"].value, cell=f"M{row}")
        s = _decimal_cell(sheet[f"N{row}"].value, cell=f"N{row}")
        h = _decimal_cell(sheet[f"O{row}"].value, cell=f"O{row}")
        formula = str(sheet[f"P{row}"].value)
        if formula == f"=2*O{row}+4*M{row}-2*N{row}":
            workbook_perimeter = 2 * h + 4 * b - 2 * s
        elif formula in {f"=M{row}*2+O{row}*2", f"=2*M{row}+2*O{row}"}:
            workbook_perimeter = 2 * b + 2 * h
        else:
            raise ObmWorkbookExportError(
                f"Row {row}: unsupported perimeter formula {formula!r}"
            )
        if abs(workbook_perimeter - heated) > Decimal("0.001"):
            raise ObmWorkbookExportError(
                f"Row {row}: perimeter conflicts with fixed workbook geometry"
            )
        count = _decimal_cell(sheet[f"G{row}"].value, cell=f"G{row}") * _decimal_cell(
            sheet[f"H{row}"].value, cell=f"H{row}"
        )
        if count != Decimal(element.quantity):  # type: ignore[arg-type]
            raise ObmWorkbookExportError(
                f"Row {row}: quantity cannot be represented without changing ambiguous G/H inputs"
            )
        if sheet[f"E{row}"].value != element.heating_sides:
            raise ObmWorkbookExportError(
                f"Row {row}: heating sides differ; workbook Q formula ignores column E"
            )


def export_obm_workbook(
    elements: Iterable[ProjectElement],
    decisions: Iterable[RequiredFireResistanceDecision],
    source_path: str | Path,
    output_path: str | Path,
    *,
    json_report: str | Path | None = None,
    markdown_report: str | Path | None = None,
    mode: ExecutionMode = ExecutionMode.DRAFT,
    technical_entry: FireproofingTechnicalEntry | None = None,
    verified_template_sha256: str | None = None,
    normative_validations: Iterable[NormativeValidation] | None = None,
    calculation_date: date | None = None,
) -> ObmWorkbookExportReport:
    """Populate only confirmed input cells in a verified copy of the fixed template."""

    if not isinstance(mode, ExecutionMode):
        raise TypeError("mode must be ExecutionMode")
    items = tuple(elements)
    choices = tuple(decisions)
    source = Path(source_path).resolve(strict=True)
    output = Path(output_path).resolve(strict=False)
    technical_status = (
        technical_entry.status
        if technical_entry is not None
        else TechnicalDataStatus.UNVERIFIED_TECHNICAL_DATA
    )
    if (
        mode is ExecutionMode.PRODUCTION
        and (
            technical_entry is None
            or calculation_date is None
            or not technical_entry.verified_for_production_on(calculation_date)
        )
    ):
        raise ObmWorkbookExportError(
            "Production Excel thickness/consumption export requires verified primary technical data"
        )
    source_hash = file_sha256(source)
    template_verified = (
        verified_template_sha256 is not None
        and source_hash == verified_template_sha256.lower()
    )
    if mode is ExecutionMode.PRODUCTION and not template_verified:
        raise ObmWorkbookExportError(
            "Production Excel export requires a verified template SHA-256"
        )
    validations = tuple(normative_validations or ())
    if mode is ExecutionMode.PRODUCTION:
        if len(validations) != len(choices) or any(
            not validation.valid_for_production
            or validation.trace != decision.normative_trace
            for decision, validation in zip(choices, validations)
        ):
            raise ObmWorkbookExportError(
                "Production Excel export requires registry/date/hash-validated NormativeTrace for every decision"
            )
    json_path = (
        Path(json_report).resolve(strict=False)
        if json_report
        else output.with_suffix(".audit.json")
    )
    md_path = (
        Path(markdown_report).resolve(strict=False)
        if markdown_report
        else output.with_suffix(".audit.md")
    )
    for target in (json_path, md_path):
        if target in {source, output}:
            raise ObmWorkbookExportError(
                f"Report path must differ from workbook paths: {target}"
            )
        if target.exists():
            raise ObmWorkbookExportError(f"Report already exists: {target}")
    workbook = load_workbook(source, data_only=False, read_only=False)
    try:
        if DATA_SHEET not in workbook.sheetnames:
            raise ObmWorkbookExportError(f"Required worksheet is missing: {DATA_SHEET}")
        formulas_before = _formula_map(workbook)
        if len(formulas_before) != EXPECTED_FORMULA_COUNT:
            raise ObmWorkbookExportError(
                f"Expected {EXPECTED_FORMULA_COUNT} formulas, found {len(formulas_before)}"
            )
        _validate_rows(items, choices, workbook)
        sheet = workbook[DATA_SHEET]
        changes: list[ExcelCellChange] = []
        fields = (
            ("mark", "C", None, lambda item: item.mark),
            ("length", "J", "m", lambda item: _quantity(item.length, Unit.METER)),
            ("area", "R", "cm2", lambda item: _quantity(item.area, Unit.SQUARE_CENTIMETER)),
            (
                "required_fire_resistance",
                "X",
                "min",
                lambda item: _quantity(item.required_fire_resistance, Unit.MINUTE),
            ),
        )
        for offset, item in enumerate(items):
            row = FIRST_ELEMENT_ROW + offset
            for field, column, unit, getter in fields:
                after = getter(item)
                before = sheet[f"{column}{row}"].value
                if str(before) != str(after):
                    changes.append(
                        ExcelCellChange(
                            f"{DATA_SHEET}!{column}{row}", field, before, after, unit
                        )
                    )
    finally:
        workbook.close()

    mapping = WorkbookMapping(
        columns=(
            ColumnBinding("mark", DATA_SHEET, "C", FIRST_ELEMENT_ROW, LAST_ELEMENT_ROW),
            ColumnBinding(
                "length",
                DATA_SHEET,
                "J",
                FIRST_ELEMENT_ROW,
                LAST_ELEMENT_ROW,
                transform=lambda value: _quantity(value, Unit.METER),
            ),
            ColumnBinding(
                "area",
                DATA_SHEET,
                "R",
                FIRST_ELEMENT_ROW,
                LAST_ELEMENT_ROW,
                transform=lambda value: _quantity(value, Unit.SQUARE_CENTIMETER),
            ),
            ColumnBinding(
                "required_fire_resistance",
                DATA_SHEET,
                "X",
                FIRST_ELEMENT_ROW,
                LAST_ELEMENT_ROW,
                transform=lambda value: _quantity(value, Unit.MINUTE),
            ),
        )
    )
    copy_result: ExcelCopyResult = write_mapped_copy(
        source, output, mapping, elements=items
    )

    try:
        with ZipFile(output) as archive:
            bad_part = archive.testzip()
    except BadZipFile as exc:
        raise ObmWorkbookExportError("Generated workbook is not a valid ZIP") from exc
    if bad_part is not None:
        raise ObmWorkbookExportError(f"Generated workbook has a corrupt part: {bad_part}")

    result_book = load_workbook(output, data_only=False, read_only=False)
    try:
        formulas_after = _formula_map(result_book)
    finally:
        result_book.close()
    if formulas_after != formulas_before:
        raise ObmWorkbookExportError("Formula map changed during Excel export")
    if file_sha256(source) != copy_result.source_sha256:
        raise ObmWorkbookExportError("Source workbook checksum changed")

    for target in (json_path, md_path):
        target.parent.mkdir(parents=True, exist_ok=True)
    warning_items = [
        "EXCEL_RECALCULATION_REQUIRED: formula caches remain stale until Microsoft Excel recalculates the copy",
        "The export is valid only for the fixed 44-row template and does not infer G/H quantity factors or heating exposure",
    ]
    if technical_status is not TechnicalDataStatus.VERIFIED_TECHNICAL_DATA:
        warning_items.insert(
            0,
            "UNVERIFIED_TECHNICAL_DATA: thickness and consumption lookup tables have no primary technical document in the workspace",
        )
    warnings = tuple(warning_items)
    report = ObmWorkbookExportReport(
        output,
        copy_result.source_sha256,
        file_sha256(output),
        len(formulas_before),
        len(formulas_after),
        formulas_after == formulas_before,
        True,
        True,
        tuple(changes),
        warnings,
        json_path,
        md_path,
        "EXCEL_COMPATIBILITY_EXPORT",
        technical_status.value,
        "VERIFIED" if template_verified else "UNVERIFIED",
        "EXCEL_RECALCULATION_REQUIRED",
    )
    json_path.write_text(
        json.dumps(report.as_dict(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    lines = [
        "# Отчёт Excel-export",
        "",
        f"- Source SHA-256: `{report.source_sha256}`",
        f"- Output SHA-256: `{report.output_sha256}`",
        f"- Формулы: {report.formula_count_before} → {report.formula_count_after}; сохранены: `{report.formulas_preserved}`",
        f"- ZIP integrity: `{report.zip_integrity}`; openpyxl: `{report.openpyxl_opened}`",
        f"- Export kind: `{report.export_kind}`",
        f"- Technical data: `{report.technical_data_status}`",
        f"- Template verification: `{report.template_verification_status}`",
        f"- Recalculation: `{report.recalculation_status}`",
        "",
        "## Изменённые входные ячейки",
        "",
    ]
    if changes:
        for change in changes:
            lines.append(
                f"- `{change.cell}` ({change.project_field}): `{change.before}` → `{change.after}` {change.unit or ''}".rstrip()
            )
    else:
        lines.append("Значения подтверждённых входных ячеек уже совпадали.")
    lines.extend(["", "## Предупреждения", ""])
    lines.extend(f"- {warning}" for warning in warnings)
    lines.append("")
    md_path.write_text("\n".join(lines), encoding="utf-8", newline="\n")
    return report
