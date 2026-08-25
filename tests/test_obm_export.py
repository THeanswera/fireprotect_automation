from datetime import datetime, timezone
from decimal import Decimal
from hashlib import sha256
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from fireprotect.decision import RequiredFireResistanceDecision
from fireprotect.excel import ObmWorkbookExportError, export_obm_workbook
from fireprotect.model import (
    ProjectElement,
    ProvenanceType,
    Quantity,
    Unit,
    ValueProvenance,
)


def _workbook(path: Path) -> None:
    workbook = Workbook()
    data = workbook.active
    data.title = "данные"
    workbook.create_sheet("Толщина ОЗ")
    workbook.create_sheet("Расход ОЗ")
    workbook.create_sheet("ТР")
    workbook.create_sheet("Лист1")
    for row in range(6, 50):
        data[f"C{row}"] = f"К{row}"
        data[f"E{row}"] = 4
        data[f"G{row}"] = 1
        data[f"H{row}"] = 1
        data[f"I{row}"] = f"=G{row}*H{row}"
        data[f"J{row}"] = 3
        data[f"L{row}"] = "I 30К1"
        data[f"M{row}"] = 299
        data[f"N{row}"] = 9
        data[f"O{row}"] = 298
        data[f"P{row}"] = f"=2*O{row}+4*M{row}-2*N{row}"
        data[f"Q{row}"] = f"=P{row}"
        data[f"R{row}"] = 110.8
        data[f"S{row}"] = f"=R{row}*100/Q{row}"
        data[f"U{row}"] = f"=MATCH(S{row},'Толщина ОЗ'!$B$3:$B$89,1)"
        data[f"V{row}"] = f"=MATCH(X{row},'Толщина ОЗ'!$D$1:$H$1,0)"
        data[f"X{row}"] = 90
        data[f"Y{row}"] = f"=INDEX('Толщина ОЗ'!$D$3:$H$89,U{row},V{row})"
        data[f"AB{row}"] = f"=I{row}*J{row}"
        data[f"AC{row}"] = f"=AB{row}*Q{row}/1000"
        data[f"AD{row}"] = f"=MATCH(S{row},'Расход ОЗ'!$B$3:$B$89,1)"
        data[f"AE{row}"] = f"=MATCH(X{row},'Расход ОЗ'!$D$1:$H$1,0)"
        data[f"AF{row}"] = f"=INDEX('Расход ОЗ'!$D$3:$H$89,AD{row},AE{row})"
        data[f"AG{row}"] = f"=AC{row}*AF{row}"
    for cell, formula in {
        "K1": "=AG50",
        "K2": "=AC50",
        "AC50": "=SUM(AC6:AC49)",
        "AG50": "=SUM(AG6:AG49)",
        "AC51": "=AC50*1.2",
        "AC52": "=AC51*1.2",
        "AG51": "=AG50*1.2",
    }.items():
        data[cell] = formula
    workbook.save(path)
    workbook.close()


def _element(row: int, *, length: str = "4") -> ProjectElement:
    values = {name: None for name in ProjectElement.field_names()}
    values.update(
        project_id="P1",
        element_id=f"E{row}",
        mark=f"К{row}",
        element_type="column",
        source_file="project.json",
        source_type="PROJECT_JSON",
        timestamp=datetime(2026, 8, 25, tzinfo=timezone.utc),
        profile_name="30К1",
        area=Quantity.of("110.8", Unit.SQUARE_CENTIMETER),
        full_perimeter=Quantity.of("1774", Unit.MILLIMETER),
        heated_perimeter=Quantity.of("1774", Unit.MILLIMETER),
        length=Quantity.of(length, Unit.METER),
        quantity=1,
        heating_sides=4,
        required_fire_resistance=Quantity.of("90", Unit.MINUTE),
    )
    values["provenance"] = {
        name: ValueProvenance(
            ProvenanceType.SOURCE, file="project.json", field=name
        )
        for name, value in values.items()
        if value is not None and name not in ProjectElement._UNTRACED_FIELDS
    }
    return ProjectElement(**values)


def _decision(confirmed: bool = True) -> RequiredFireResistanceDecision:
    return RequiredFireResistanceDecision(
        Quantity.of("90", Unit.MINUTE),
        "column",
        "II",
        "engineer assignment; normative source pending",
        None,
        "explicit engineer input",
        confirmed,
    )


def test_fixed_obm_export_preserves_all_formulas_and_reports_changes(tmp_path: Path):
    source = tmp_path / "template.xlsx"
    output = tmp_path / "result.xlsx"
    _workbook(source)
    source_hash = sha256(source.read_bytes()).hexdigest()
    elements = [_element(row) for row in range(6, 50)]

    report = export_obm_workbook(
        elements, [_decision() for _ in elements], source, output
    )

    assert sha256(source.read_bytes()).hexdigest() == source_hash
    assert report.formula_count_before == report.formula_count_after == 579
    assert report.formulas_preserved and report.zip_integrity and report.openpyxl_opened
    assert len([item for item in report.cell_changes if item.project_field == "length"]) == 44
    assert "UNVERIFIED_TECHNICAL_DATA" in " ".join(report.warnings)
    assert report.json_report.exists() and report.markdown_report.exists()
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook["данные"]["J6"].value == 4
        assert workbook["данные"]["P6"].value == "=2*O6+4*M6-2*N6"
    finally:
        workbook.close()


def test_obm_export_blocks_unconfirmed_r_and_wrong_fixed_row_count(tmp_path: Path):
    source = tmp_path / "template.xlsx"
    _workbook(source)
    elements = [_element(row) for row in range(6, 50)]
    decisions = [_decision() for _ in elements]
    decisions[0] = _decision(False)

    with pytest.raises(ValueError, match="engineer confirmation"):
        export_obm_workbook(elements, decisions, source, tmp_path / "blocked.xlsx")
    with pytest.raises(ObmWorkbookExportError, match="44 calculation rows"):
        export_obm_workbook(
            elements[:-1], decisions[:-1], source, tmp_path / "short.xlsx"
        )
