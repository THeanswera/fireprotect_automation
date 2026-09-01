from hashlib import sha256
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from fireprotect.excel.obm import ObmWorkbookExportError, export_obm_workbook
from fireprotect.excel.registry import (
    ExcelTemplateEntry,
    ExcelTemplateRegistry,
    ExcelTemplateStatus,
    formula_map_fingerprint,
)
from fireprotect.excel.writer import file_sha256
from fireprotect.execution import ExecutionMode
from tests.test_fireproofing_registry import write_technical_registry


def test_unverified_technical_data_blocks_before_excel_copy(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"not opened because production gate must fail first")
    before = sha256(source.read_bytes()).hexdigest()
    entry = write_technical_registry(tmp_path, verified=False).entries["TEST_SYSTEM"]

    with pytest.raises(ObmWorkbookExportError, match="technical data"):
        export_obm_workbook(
            (),
            (),
            source,
            tmp_path / "output.xlsx",
            mode=ExecutionMode.PRODUCTION,
            technical_entry=entry,
            calculation_date=date(2026, 8, 25),
        )
    assert sha256(source.read_bytes()).hexdigest() == before
    assert not (tmp_path / "output.xlsx").exists()


def test_lookup_table_fingerprint_mismatch_blocks_production(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    data = workbook.active
    data.title = "данные"
    for index in range(579):
        data.cell(row=index // 100 + 1, column=index % 100 + 1).value = "=1"
    workbook.create_sheet("Толщина ОЗ")["A1"] = "controlled thickness"
    workbook.create_sheet("Расход ОЗ")["A1"] = "controlled consumption"
    workbook.save(source)
    workbook.close()

    formula_book = load_workbook(source, data_only=False)
    try:
        formulas = {
            f"{sheet.title}!{cell.coordinate}": cell.value
            for sheet in formula_book.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        }
    finally:
        formula_book.close()
    technical_entry = write_technical_registry(
        tmp_path, verified=True
    ).entries["TEST_SYSTEM"]
    template_entry = ExcelTemplateEntry(
        template_id="TEST_TEMPLATE",
        sha256=file_sha256(source),
        workbook_identity="controlled workbook",
        workbook_version="1",
        expected_formula_count=579,
        formula_map_sha256=formula_map_fingerprint(formulas),
        status=ExcelTemplateStatus.APPROVED,
        approved_by="test reviewer",
        approved_at=date(2026, 8, 25),
        technical_data_entry_id=technical_entry.entry_id,
        technical_data_version=technical_entry.version,
        lookup_sheets=("Толщина ОЗ", "Расход ОЗ"),
        lookup_table_sha256="0" * 64,
    )

    with pytest.raises(ObmWorkbookExportError, match="EXCEL_LOOKUP_TABLE_MISMATCH"):
        export_obm_workbook(
            (),
            (),
            source,
            tmp_path / "output.xlsx",
            mode=ExecutionMode.PRODUCTION,
            technical_entry=technical_entry,
            template_entry=template_entry,
            calculation_date=date(2026, 8, 25),
        )
    assert not (tmp_path / "output.xlsx").exists()


def test_workspace_obm_registry_entry_remains_unverified():
    registry_path = Path(__file__).resolve().parents[1] / "templates" / "excel_registry.yaml"
    entry = ExcelTemplateRegistry.load(registry_path).require(
        "OBM_WORKBOOK_UNVERIFIED"
    )
    assert entry.status is ExcelTemplateStatus.UNVERIFIED
    assert entry.approved_by is None


def test_run_cannot_self_certify_excel_template_hash(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"controlled placeholder")
    entry = write_technical_registry(tmp_path, verified=True).entries["TEST_SYSTEM"]

    template_entry = ExcelTemplateEntry(
        template_id="TEST_TEMPLATE",
        sha256="0" * 64,
        workbook_identity="controlled workbook",
        workbook_version="1",
        expected_formula_count=0,
        formula_map_sha256="0" * 64,
        status=ExcelTemplateStatus.APPROVED,
        approved_by="test reviewer",
        approved_at=date(2026, 8, 25),
        technical_data_entry_id=entry.entry_id,
        technical_data_version=entry.version,
        lookup_sheets=("lookup",),
        lookup_table_sha256="0" * 64,
    )

    with pytest.raises(ObmWorkbookExportError, match="trusted template_id"):
        export_obm_workbook(
            (),
            (),
            source,
            tmp_path / "output.xlsx",
            mode=ExecutionMode.PRODUCTION,
            technical_entry=entry,
            template_entry=template_entry,
            calculation_date=date(2026, 8, 25),
        )
    assert not (tmp_path / "output.xlsx").exists()
