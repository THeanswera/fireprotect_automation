"""The review workbook is new, self-describing and free of invented results."""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from fireprotect.excel.review import (
    ReviewWorkbookError,
    export_lira_bar_review,
)


def _run_manifest(tmp_path: Path, *, prepared: bool = False) -> Path:
    run_dir = tmp_path / "run"
    (run_dir / "experiment").mkdir(parents=True)
    (run_dir / "experiment" / "experiment_input.json").write_text(
        json.dumps(
            {
                "kind": "LIRA_BAR_EXPERIMENT_INPUT",
                "status": "EXPERIMENT_INPUT_READY",
                "selected_row": {
                    "row_id": "R0003",
                    "element_id": "1",
                    "section_station": "2",
                    "rsu_group": "A1",
                    "rsu_criterion": "1",
                    "rsu_column_number": 2,
                    "load_case_membership": ["1", "2", "3"],
                    "selection_basis": "учебная строка",
                    "row": {
                        "source_terms": [
                            {
                                "load_case_id": "1",
                                "coefficient": "1.0",
                                "forces": {
                                    "N": {"value": "0.0"},
                                    "Mk": {"value": "0.0"},
                                    "My": {"value": "0.1125"},
                                    "Mz": {"value": "0.0"},
                                    "Qy": {"value": "0.0"},
                                    "Qz": {"value": "0.0"},
                                },
                                "coefficient_source": {"cell": "E4"},
                            }
                        ]
                    },
                },
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    manifest = {
        "kind": "LIRA_BAR_RUN_MANIFEST",
        "status": "RUN_PREPARED",
        "run_id": "RUN_TEST",
        "source_package": "source",
        "experiment_row_id": "R0003",
        "plan_reference": {"path": "PLAN.md", "sha256": "a" * 64},
        "engineer_confirmed": False,
        "declaration_status": "DRAFT_UNSIGNED",
        "bar": {
            "bar_id": "Б2",
            "element_ids": ["1", "2"],
            "end_node_ids": ["1", "2"],
            "element_lengths_m": {"1": "1.5", "2": "1.5"},
            "bar_length_m": "3.0",
            "join_basis": "одна прямая цепочка КЭ",
        },
        "profile": {
            "kind_word": "Швеллер",
            "designation": "22П",
            "mark": "Б2",
            "standard": "ГОСТ 8240-97",
            "rotation_degrees": "0",
            "plane": "X-Z",
            "scheme_flag": "2",
            "rx3_template": "Б2",
            "stress_state": "ONE_PLANE_BENDING",
        },
        "design_conditions": {
            "effective_length_m": None,
            "support_condition": None,
            "heating_sides": "4",
            "fire_regime": "STANDARD",
            "required_fire_resistance_min": "R15",
        },
        "decisions": {
            "design_conditions": {
                "role": "ASSISTANT_SELECTION",
                "basis": "учебная постановка",
                "reference": None,
            }
        },
        "components": [
            {
                "component": "My",
                "source_value": "0.63",
                "source_unit": "tf*m",
                "source_cell": "I6",
                "review_value": "6.1781895",
                "review_unit": "kN*m",
                "convention": {
                    "resolved": True,
                    "target": "FIELD50_MAX_MAJOR_AXIS_MOMENT",
                    "value_transform": "MAGNITUDE",
                },
            }
        ],
        "source_files": {
            "published": {
                "path": "РСУ.xls",
                "recorded_sha256": "b" * 64,
                "current_sha256": "c" * 64,
            }
        },
        "blockers": [],
        "missing_confirmations": ["effective_length_m", "support_condition"],
        "governing_result_selection": None,
        "rx38_created": False,
        "release_forbidden": True,
        "issue_readiness": "NOT_READY_FOR_ISSUE",
    }
    if prepared:
        (run_dir / "rx3_input").mkdir()
        (run_dir / "rx3_input" / "rx3_validation_manifest.json").write_text(
            json.dumps(
                {
                    "generated": {"path": "generated.rx38", "sha256": "d" * 64},
                    "target": {"position": 12, "after_fingerprint": "e" * 64},
                    "effective_length_and_support": {
                        "status": "NOT_APPLICABLE_FOR_THIS_RECORD",
                        "basis": ["выбранная строка не содержит осевого усилия"],
                        "not_substituted": "геометрические 3 м не назначались",
                    },
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
    path = run_dir / "run_manifest.json"
    path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return path


def test_writes_a_new_self_describing_workbook(tmp_path: Path) -> None:
    manifest = _run_manifest(tmp_path, prepared=True)
    output = tmp_path / "review.xlsx"

    result = export_lira_bar_review(run_manifest=manifest, output=output)

    assert result["status"] == "REVIEW_WORKBOOK_WRITTEN"
    assert result["calculation_results_included"] is False
    assert result["release_forbidden"] is True
    workbook = load_workbook(output)
    assert workbook.sheetnames == [
        "Статус",
        "Идентичность",
        "Строка РСУ",
        "Усилия",
        "Условия",
        "Источники",
        "Блокеры",
        "Результат RX3",
    ]
    status = {row[0].value: row[1].value for row in workbook["Статус"].iter_rows(min_row=2)}
    assert "НЕ ДЛЯ ВЫПУСКА" in str(status["Статус файла"])
    assert status["Строка технического опыта"] == "R0003"
    assert status["Подпись инженера"].startswith("НЕТ")
    assert status["Файл RX3 для ручного расчёта"] == "generated.rx38"

    identity = {
        row[0].value: row[1].value for row in workbook["Идентичность"].iter_rows(min_row=2)
    }
    assert identity["Профиль"] == "22П"
    assert identity["Геометрическая длина, м"] == "3.0"

    forces = list(workbook["Усилия"].iter_rows(min_row=2, values_only=True))
    assert forces[0][0] == "My"
    assert forces[0][4] == "6.1781895"
    assert forces[0][6] == "FIELD50_MAX_MAJOR_AXIS_MOMENT"

    row_sheet = list(workbook["Строка РСУ"].iter_rows(values_only=True))
    assert any(row[0] == "row_id" and row[1] == "R0003" for row in row_sheet)
    assert any(row[0] == "1" and row[1] == "1.0" for row in row_sheet)

    conditions = list(workbook["Условия"].iter_rows(values_only=True))
    assert any(
        row[0] == "effective_length_m" and row[1] == "не задано"
        for row in conditions
    )

    sources = list(workbook["Источники"].iter_rows(min_row=2, values_only=True))
    assert sources[0][4] == "НЕТ"


def test_result_cells_stay_empty_before_a_real_calculation(tmp_path: Path) -> None:
    manifest = _run_manifest(tmp_path)
    output = tmp_path / "review.xlsx"

    export_lira_bar_review(run_manifest=manifest, output=output)

    workbook = load_workbook(output)
    rows = list(workbook["Результат RX3"].iter_rows(min_row=2, values_only=True))
    results = {row[0]: row[1] for row in rows if row[0]}
    assert results["Критическая температура, °C"] is None
    assert results["Собственный предел огнестойкости, мин"] is None
    assert results["Коэффициент использования по моменту"] is None
    assert results["Требуемая толщина огнезащиты, мм"] is None
    assert results["Расход материала"] is None
    assert "не заполнено" in str(
        {row[0]: row[2] for row in rows if row[0]}[
            "Критическая температура, °C"
        ]
    )


def test_existing_workbook_is_never_overwritten(tmp_path: Path) -> None:
    manifest = _run_manifest(tmp_path)
    output = tmp_path / "review.xlsx"
    output.write_bytes(b"existing")
    with pytest.raises(ReviewWorkbookError, match="refusing to overwrite"):
        export_lira_bar_review(run_manifest=manifest, output=output)
    assert output.read_bytes() == b"existing"


def test_foreign_manifest_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "run_manifest.json"
    path.write_text(json.dumps({"kind": "SOMETHING_ELSE"}), encoding="utf-8")
    with pytest.raises(ReviewWorkbookError, match="LIRA_BAR_RUN_MANIFEST"):
        export_lira_bar_review(run_manifest=path, output=tmp_path / "review.xlsx")


def test_output_must_be_xlsx(tmp_path: Path) -> None:
    manifest = _run_manifest(tmp_path)
    with pytest.raises(ReviewWorkbookError, match="must be an .xlsx"):
        export_lira_bar_review(run_manifest=manifest, output=tmp_path / "review.csv")
