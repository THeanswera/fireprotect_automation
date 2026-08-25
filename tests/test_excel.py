from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from fireprotect.excel import (
    CellBinding,
    ColumnBinding,
    CopyOnlyViolationError,
    FormulaOverwriteError,
    WorkbookMapping,
    WorkbookMappingError,
    file_sha256,
    write_mapped_copy,
    write_project_elements_copy,
)
from fireprotect.model import ProjectElement, ProvenanceType, ValueProvenance


MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
M = f"{{{MAIN_NS}}}"


@dataclass
class Item:
    amount: Decimal


def _make_workbook(path: Path) -> None:
    parts = {
        "[Content_Types].xml": """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
</Types>""",
        "_rels/.rels": """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>""",
        "xl/workbook.xml": """<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
 xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets><sheet name="data" sheetId="1" r:id="rId1"/></sheets>
  <calcPr calcId="191029"/>
</workbook>""",
        "xl/_rels/workbook.xml.rels": """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>""",
        "xl/worksheets/sheet1.xml": """<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <dimension ref="A1:C3"/>
  <sheetData>
    <row r="1"><c r="A1" s="1" t="inlineStr"><is><t>Template</t></is></c></row>
    <row r="2">
      <c r="A2" s="7"><v>10</v></c>
      <c r="B2" s="9"><f>A2*2</f><v>20</v></c>
      <c r="C2" s="5"/>
    </row>
    <row r="3">
      <c r="A3" s="7"><v>11</v></c>
      <c r="B3" s="9"><f>A3*2</f><v>22</v></c>
      <c r="C3" s="5"/>
    </row>
  </sheetData>
</worksheet>""",
        "xl/styles.xml": """<?xml version="1.0" encoding="UTF-8"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="1"><font><name val="Calibri"/><sz val="11"/></font></fonts>
  <fills count="2"><fill><patternFill/></fill><fill><patternFill patternType="gray125"/></fill></fills>
  <borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="10">
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
    <xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>
  </cellXfs>
  <cellStyles count="1"><cellStyle name="Normal" xfId="0" builtinId="0"/></cellStyles>
  <tableStyles count="0" defaultTableStyle="TableStyleMedium9" defaultPivotStyle="PivotStyleLight16"/>
</styleSheet>""",
    }
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as workbook:
        for name, xml in parts.items():
            workbook.writestr(name, xml)


def _sheet_root(path: Path) -> ET.Element:
    with ZipFile(path) as workbook:
        return ET.fromstring(workbook.read("xl/worksheets/sheet1.xml"))


def _cell(root: ET.Element, coordinate: str) -> ET.Element:
    return next(cell for cell in root.iter(M + "c") if cell.get("r") == coordinate)


def test_writes_configured_cells_and_columns_to_copy_only(tmp_path: Path) -> None:
    source = tmp_path / "template.xlsx"
    output = tmp_path / "result.xlsx"
    _make_workbook(source)
    checksum = file_sha256(source)
    with ZipFile(source) as workbook:
        original_styles = workbook.read("xl/styles.xml")

    mapping = WorkbookMapping(
        cells=(CellBinding("summary.name", "data", "C2"),),
        columns=(
            ColumnBinding("amount", "data", "A", first_row=2, last_row=3),
        ),
    )
    result = write_mapped_copy(
        source,
        output,
        mapping,
        element={"summary": {"name": "  calculation  "}},
        elements=(Item(Decimal("12.50")), Item(Decimal("3"))),
    )

    assert source.exists()
    assert file_sha256(source) == checksum == result.source_sha256
    assert output.exists() and output != source
    assert result.written_cells == ("data!C2", "data!A2", "data!A3")

    root = _sheet_root(output)
    assert _cell(root, "A2").find(M + "v").text == "12.50"
    assert _cell(root, "A3").find(M + "v").text == "3"
    written_text = _cell(root, "C2").find(f"{M}is/{M}t")
    assert written_text.text == "  calculation  "
    assert written_text.get("{http://www.w3.org/XML/1998/namespace}space") == "preserve"

    # Existing styles and all non-target formulas survive the export.
    assert _cell(root, "A2").get("s") == "7"
    assert _cell(root, "C2").get("s") == "5"
    assert _cell(root, "B2").get("s") == "9"
    assert _cell(root, "B2").find(M + "f").text == "A2*2"
    assert _cell(root, "B3").find(M + "f").text == "A3*2"
    with ZipFile(output) as workbook:
        assert workbook.read("xl/styles.xml") == original_styles
        workbook_xml = ET.fromstring(workbook.read("xl/workbook.xml"))
    calc = workbook_xml.find(M + "calcPr")
    assert calc.get("fullCalcOnLoad") == "1"
    assert calc.get("forceFullCalc") == "1"
    assert calc.get("calcMode") == "auto"

    # A real Excel reader accepts the produced OOXML and sees formulas/styles.
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        assert workbook["data"]["B2"].value == "=A2*2"
        assert workbook["data"]["B2"].style_id == 9
        assert workbook.calculation.fullCalcOnLoad is True
        assert workbook.calculation.forceFullCalc is True
    finally:
        workbook.close()


def test_empty_mapping_is_a_binary_copy(tmp_path: Path) -> None:
    source = tmp_path / "template.xlsx"
    output = tmp_path / "result.xlsx"
    _make_workbook(source)

    result = write_mapped_copy(source, output, WorkbookMapping())

    assert file_sha256(output) == file_sha256(source) == result.source_sha256
    assert result.written_cells == ()


def test_rejects_source_as_output_and_existing_output(tmp_path: Path) -> None:
    source = tmp_path / "template.xlsx"
    existing = tmp_path / "existing.xlsx"
    _make_workbook(source)
    _make_workbook(existing)

    with pytest.raises(CopyOnlyViolationError):
        write_mapped_copy(source, source, WorkbookMapping())
    with pytest.raises(CopyOnlyViolationError):
        write_mapped_copy(source, existing, WorkbookMapping())


def test_formula_target_is_protected_and_no_output_is_left(tmp_path: Path) -> None:
    source = tmp_path / "template.xlsx"
    output = tmp_path / "result.xlsx"
    _make_workbook(source)
    checksum = file_sha256(source)
    mapping = WorkbookMapping(cells=(CellBinding("value", "data", "B2"),))

    with pytest.raises(FormulaOverwriteError):
        write_mapped_copy(source, output, mapping, element={"value": 99})

    assert not output.exists()
    assert file_sha256(source) == checksum


def test_requires_existing_cells_and_respects_column_capacity(tmp_path: Path) -> None:
    source = tmp_path / "template.xlsx"
    _make_workbook(source)

    with pytest.raises(WorkbookMappingError, match="does not exist"):
        write_mapped_copy(
            source,
            tmp_path / "missing.xlsx",
            WorkbookMapping(cells=(CellBinding("value", "data", "D8"),)),
            element={"value": 1},
        )
    with pytest.raises(WorkbookMappingError, match="accepts 1 rows"):
        write_mapped_copy(
            source,
            tmp_path / "overflow.xlsx",
            WorkbookMapping(
                columns=(ColumnBinding("amount", "data", "A", 2, 2),)
            ),
            elements=(Item(Decimal("1")), Item(Decimal("2"))),
        )


def test_project_element_list_is_written_through_explicit_mapping(tmp_path: Path) -> None:
    source = tmp_path / "template.xlsx"
    output = tmp_path / "result.xlsx"
    _make_workbook(source)
    values = {name: None for name in ProjectElement.field_names()}
    values.update(
        project_id="P1",
        element_id="E1",
        mark="К1",
        element_type="column",
        source_file="project.json",
        source_type="PROJECT_JSON",
        timestamp=datetime(2026, 8, 25, tzinfo=timezone.utc),
    )
    values["provenance"] = {
        name: ValueProvenance(
            ProvenanceType.SOURCE, file="project.json", field=name
        )
        for name in ("project_id", "element_id", "mark", "element_type")
    }
    element = ProjectElement(**values)

    result = write_project_elements_copy(
        source,
        output,
        [element],
        WorkbookMapping(columns=(ColumnBinding("mark", "data", "C", 2, 3),)),
    )

    assert result.written_cells == ("data!C2",)
    workbook = load_workbook(output, data_only=False)
    try:
        assert workbook["data"]["C2"].value == "К1"
    finally:
        workbook.close()
